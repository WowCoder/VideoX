"""脚本路由：CRUD + 分镜单元 + 批量更新 + 脚本解析 + 初剪。"""

import csv
import io
import json
import os
import textwrap
import uuid
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import delete, desc, exists, func, or_, select, update
from sqlalchemy.orm import selectinload

from app.database import get_db
from app.dependencies import get_current_user
from app.models.user import User
from app.models.project import Project
from app.models.script import Script, ScriptUnit, Shot, Candidate, ShotGeneration
from app.models.cut_version import CutVersion
from app.models.job import Job
from app.models.asset import Asset, Clip
from app.services.project_access import accessible_project_condition, require_project_access
from app.schemas.script import (
    ScriptCreate, ScriptUpdate, ScriptOut, ScriptListOut,
    ScriptUnitBatchUpdate, ScriptEditorSave, LinkedCutOut, ShotOut,
)
from app.utils.script_parser import parse_script_text
from app.utils.duration import format_duration
from app.tasks.celery_app import celery_app
from app.config import get_settings

router = APIRouter()
settings = get_settings()


class MatchRequest(BaseModel):
    unit_id: str | None = None
    shot_ids: list[str] = Field(default_factory=list, max_length=200)


class RoughCutRequest(BaseModel):
    title: str | None = None
    visual_config: dict | None = None


class CompleteShotsRequest(BaseModel):
    shot_ids: list[str] = Field(..., min_length=1, max_length=50)


def _compute_metrics(units: list[ScriptUnit]) -> dict:
    """按分镜统计时长和画面完成度，避免把内部 ScriptUnit 暴露为进度口径。"""
    total_seconds = 0
    matched_count = 0
    shot_count = 0
    for unit in units:
        if unit.shots:
            for shot in unit.shots:
                shot_count += 1
                total_seconds += shot.duration or 8
                if (
                    shot.selected_clip_id
                    or shot.generated_image_path
                    or shot.generated_video_path
                    or shot.generated_effect_video_path
                ):
                    matched_count += 1
        else:
            shot_count += 1
            total_seconds += unit.duration or 8
            if unit.selected_clip_id:
                matched_count += 1

    return {
        "duration": format_duration(total_seconds),
        "matched_count": matched_count,
        "progress": round(matched_count / shot_count * 100) if shot_count else 0,
        "unit_count": shot_count,
    }


def _script_export_data(script: Script) -> dict:
    return {
        "id": script.id,
        "project_id": script.project_id,
        "title": script.title,
        "summary": script.summary,
        "status": script.status,
        "duration": script.duration,
        "format": script.format,
        "updated_at": script.updated_at.isoformat() if script.updated_at else None,
        "units": [
            {
                "index": unit.index,
                "outline": unit.outline,
                "text": unit.text,
                "intent": unit.intent,
                "status": unit.status,
                "duration": unit.duration,
                "notes": unit.notes,
                "selected_clip_id": unit.selected_clip_id,
                "shots": [
                    {
                        "index": shot.index,
                        "description": shot.description,
                        "status": shot.status,
                        "visual_mode": shot.visual_mode,
                        "duration": shot.duration,
                        "notes": shot.notes,
                        "selected_clip_id": shot.selected_clip_id,
                    }
                    for shot in unit.shots
                ],
            }
            for unit in script.units
        ],
    }


def _script_rows(script: Script) -> list[list[object]]:
    rows: list[list[object]] = []
    for unit in script.units:
        shots = unit.shots or [None]
        for shot in shots:
            rows.append([
                unit.index + 1,
                unit.outline,
                unit.text,
                unit.intent,
                unit.status,
                unit.duration,
                "" if shot is None else shot.index + 1,
                "" if shot is None else shot.description,
                "" if shot is None else (shot.visual_mode or ""),
                "" if shot is None else shot.status,
                "" if shot is None else shot.duration,
                unit.notes if shot is None else shot.notes,
            ])
    return rows


def _pdf_font_path() -> str:
    candidates = [
        settings.subtitle_font_path,
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    raise RuntimeError("服务器缺少可用于脚本 PDF 的字体")


def _script_pdf(script: Script) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    font_path = _pdf_font_path()
    title_font = ImageFont.truetype(font_path, 34)
    body_font = ImageFont.truetype(font_path, 22)
    lines = [script.title, f"状态：{script.status}    时长：{script.duration}    画幅：{script.format}"]
    if script.summary:
        lines.extend(textwrap.wrap(f"摘要：{script.summary}", width=48))
    lines.append("")
    for unit in script.units:
        lines.extend(textwrap.wrap(f"{unit.index + 1}. {unit.outline or '未命名单元'}", width=48))
        if unit.text:
            lines.extend(textwrap.wrap(f"台词：{unit.text}", width=48))
        if unit.intent:
            lines.extend(textwrap.wrap(f"画面意图：{unit.intent}", width=48))
        for shot in unit.shots:
            lines.extend(
                textwrap.wrap(
                    f"  镜头 {shot.index + 1}：{shot.description}（{shot.duration}s，{shot.status}）",
                    width=46,
                )
            )
        lines.append("")

    page_width, page_height = 1240, 1754
    margin, line_height = 80, 34
    max_lines = (page_height - margin * 2) // line_height
    pages = []
    for page_index in range(0, len(lines), max_lines):
        page = Image.new("RGB", (page_width, page_height), "white")
        draw = ImageDraw.Draw(page)
        y = margin
        for line_index, line in enumerate(lines[page_index:page_index + max_lines]):
            font = title_font if page_index == 0 and line_index == 0 else body_font
            draw.text((margin, y), line, fill="#182033", font=font)
            y += 48 if font is title_font else line_height
        pages.append(page)

    output = io.BytesIO()
    pages[0].save(output, format="PDF", save_all=True, append_images=pages[1:], resolution=144)
    return output.getvalue()


@router.get("", response_model=ScriptListOut)
async def list_scripts(
    project_id: str = Query("", description="按项目筛选"),
    search: str = Query("", description="搜索关键词"),
    status_filter: str = Query("", alias="status"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """获取脚本列表。"""
    query = select(Script).join(Project, Script.project_id == Project.id).where(
        accessible_project_condition(user.id),
        Script.is_snapshot.is_(False),
    )

    if project_id:
        query = query.where(Script.project_id == project_id)
    if search:
        keyword = f"%{search.strip()}%"
        unit_match = exists(
            select(1).where(
                ScriptUnit.script_id == Script.id,
                or_(
                    ScriptUnit.outline.ilike(keyword),
                    ScriptUnit.text.ilike(keyword),
                    ScriptUnit.intent.ilike(keyword),
                    exists(
                        select(1).where(
                            Shot.unit_id == ScriptUnit.id,
                            Shot.description.ilike(keyword),
                        )
                    ),
                ),
            )
        )
        query = query.where(
            or_(Script.title.ilike(keyword), Script.summary.ilike(keyword), unit_match)
        )
    if status_filter:
        query = query.where(Script.status == status_filter)

    count_query = select(func.count()).select_from(query.subquery())
    total = (await db.execute(count_query)).scalar() or 0

    query = query.options(
        selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations)
    ).order_by(desc(Script.created_at)).offset((page - 1) * page_size).limit(page_size)
    result = await db.execute(query)
    scripts = result.scalars().all()

    return ScriptListOut(items=[ScriptOut.model_validate(s) for s in scripts], total=total)


@router.post("", response_model=ScriptOut, status_code=201)
async def create_script(
    req: ScriptCreate,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """创建脚本。

    支持 source_type: blank / paste / spreadsheet / ai
    当 source_type 为 paste/spreadsheet 且提供了文本时，自动解析。
    """
    await require_project_access(req.project_id, user, db, "editor")

    # 创建脚本
    visual_config = req.visual_config or {
        "stylePreset": "brand-tech",
        "aspectRatio": "16:9",
        "palette": "冷灰蓝",
        "cameraRhythm": "稳定推进",
        "consistencyEnabled": True,
        "generationPolicy": "优先使用项目素材，缺失镜头由 AI 补齐",
        "references": {"characters": [], "objects": [], "spaces": []},
    }

    script = Script(
        project_id=req.project_id,
        title=req.title,
        summary=req.summary,
        format=req.format,
        source_type=req.source_type,
        visual_config=visual_config,
        status="编辑中",
        needs_initial_save=True,
    )
    db.add(script)
    await db.flush()

    # 创建分镜单元
    for idx, unit_data in enumerate(req.units):
        unit = ScriptUnit(
            script_id=script.id,
            index=idx,
            outline=unit_data.outline or (unit_data.text[:16] if unit_data.text else ""),
            text=unit_data.text,
            intent=unit_data.intent or (unit_data.text[:30] if unit_data.text else ""),
            duration=unit_data.duration or 8,
            notes=unit_data.notes,
            status="待编辑",
        )
        db.add(unit)
        await db.flush()

        # 创建默认 shot
        shot = Shot(
            unit_id=unit.id,
            index=0,
            description=unit_data.intent or (unit_data.text[:30] if unit_data.text else ""),
            duration=unit_data.duration or 8,
            notes=unit_data.notes,
            status="待编辑",
        )
        db.add(shot)

    await db.flush()

    # 重算指标（预加载 shots 避免异步懒加载错误）
    units_result = await db.execute(
        select(ScriptUnit)
        .where(ScriptUnit.script_id == script.id)
        .options(selectinload(ScriptUnit.shots).selectinload(Shot.generations))
        .order_by(ScriptUnit.index)
    )
    units = list(units_result.scalars().all())
    metrics = _compute_metrics(units)
    script.duration = metrics["duration"]
    script.unit_count = metrics["unit_count"]
    script.matched_count = metrics["matched_count"]
    script.progress = metrics["progress"]

    # 更新项目
    proj_result = await db.execute(select(Project).where(Project.id == req.project_id))
    project = proj_result.scalar_one_or_none()
    if project:
        project.script_count = (project.script_count or 0) + 1
        project.status = "脚本编辑中"

    await db.flush()

    # 重新加载脚本（含 units 和 shots），避免异步序列化错误
    fresh_result = await db.execute(
        select(Script)
        .where(Script.id == script.id)
        .options(selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations))
    )
    script = fresh_result.scalar_one()
    return ScriptOut.model_validate(script)


@router.post("/parse", response_model=list[dict])
async def parse_script(
    text: str = "",
    source_type: str = "paste",
    user: User = Depends(get_current_user),
):
    """解析脚本文本，返回预览单元列表（不保存）。"""
    units = parse_script_text(text, source_type)
    return units


@router.get("/snapshot", response_model=ScriptOut)
async def get_project_script_snapshot(
    project_id: str = Query(..., min_length=1),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """获取剪辑项目创建初剪时保存的脚本基线。"""
    script = (
        await db.execute(
            select(Script)
            .join(Project, Script.project_id == Project.id)
            .where(
                Script.project_id == project_id,
                Script.is_snapshot.is_(True),
                accessible_project_condition(user.id),
            )
            .options(selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations))
            .order_by(desc(Script.created_at))
        )
    ).scalars().first()
    if not script:
        raise HTTPException(status_code=404, detail="初剪脚本快照不存在")
    return ScriptOut.model_validate(script)


@router.get("/{script_id}", response_model=ScriptOut)
async def get_script(
    script_id: str,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """获取脚本详情（含分镜单元和 shots）。"""
    result = await db.execute(
        select(Script)
        .join(Project, Script.project_id == Project.id)
        .where(Script.id == script_id, accessible_project_condition(user.id))
        .options(
            selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations),
        )
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    return ScriptOut.model_validate(script)


@router.get("/{script_id}/cuts", response_model=list[LinkedCutOut])
async def list_linked_cuts(
    script_id: str,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """列出由当前原脚本创建、且当前用户仍可访问的剪辑版本。"""
    script = (
        await db.execute(
            select(Script)
            .join(Project, Script.project_id == Project.id)
            .where(Script.id == script_id, accessible_project_condition(user.id))
        )
    ).scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")

    projects = (
        (
            await db.execute(
                select(Project)
                .where(
                    Project.source_script_id == script_id,
                    accessible_project_condition(user.id),
                )
                .order_by(desc(Project.updated_at))
            )
        ).scalars().all()
    )
    result: list[LinkedCutOut] = []
    for project in projects:
        access = await require_project_access(project.id, user, db)
        versions = (
            await db.execute(
                select(CutVersion)
                .where(CutVersion.project_id == project.id)
                .order_by(desc(CutVersion.created_at))
            )
        ).scalars().all()
        for version in versions:
            result.append(
                LinkedCutOut(
                    project_id=project.id,
                    version_id=version.id,
                    name=project.name,
                    version_name=version.name,
                    status=version.status,
                    aspect_ratio=version.aspect_ratio,
                    timeline_version=version.timeline_version or 0,
                    current_user_role=access.role,
                    updated_at=version.updated_at,
                )
            )
    return result


@router.get("/{script_id}/export")
async def export_script(
    script_id: str,
    format: str = Query("json", pattern="^(json|csv|xlsx|pdf)$"),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """把当前数据库中的真实脚本与分镜导出为 JSON/CSV/XLSX/PDF。"""
    result = await db.execute(
        select(Script)
        .join(Project, Script.project_id == Project.id)
        .where(Script.id == script_id, accessible_project_condition(user.id))
        .options(selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations))
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")

    headers = [
        "单元序号", "大纲", "台词/旁白", "画面意图", "单元状态", "单元时长(秒)",
        "镜头序号", "镜头描述", "视觉模式", "镜头状态", "镜头时长(秒)", "备注",
    ]
    media_type = "application/octet-stream"
    if format == "json":
        content = json.dumps(_script_export_data(script), ensure_ascii=False, indent=2).encode("utf-8")
        media_type = "application/json; charset=utf-8"
    elif format == "csv":
        text_output = io.StringIO()
        writer = csv.writer(text_output)
        writer.writerow(headers)
        writer.writerows(_script_rows(script))
        content = ("\ufeff" + text_output.getvalue()).encode("utf-8")
        media_type = "text/csv; charset=utf-8"
    elif format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "分镜脚本"
        sheet.append(headers)
        for row in _script_rows(script):
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        widths = [10, 22, 40, 36, 12, 14, 10, 40, 12, 12, 14, 28]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        binary = io.BytesIO()
        workbook.save(binary)
        content = binary.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        try:
            content = _script_pdf(script)
        except RuntimeError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error
        media_type = "application/pdf"

    filename = quote(f"{script.title}.{format}")
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.patch("/{script_id}", response_model=ScriptOut)
async def update_script(
    script_id: str,
    req: ScriptUpdate,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """更新脚本基本信息。"""
    result = await db.execute(
        select(Script).join(Project, Script.project_id == Project.id).where(
            Script.id == script_id, accessible_project_condition(user.id)
        )
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    await require_project_access(script.project_id, user, db, "editor")

    update_data = req.model_dump(exclude_unset=True)
    if "visual_config" in update_data and update_data["visual_config"]:
        references = (update_data["visual_config"].get("references") or {})
        referenced_assets: set[str] = set()
        referenced_clips: dict[str, str] = {}
        for category in ("characters", "objects", "spaces"):
            for reference in references.get(category, []):
                asset_id = reference.get("assetId")
                clip_id = reference.get("clipId")
                if reference.get("category") and reference["category"] != category:
                    raise HTTPException(status_code=400, detail="统一素材引用分类不一致")
                if clip_id and not asset_id:
                    raise HTTPException(status_code=400, detail="片段引用必须同时包含素材 ID")
                if asset_id:
                    referenced_assets.add(asset_id)
                if clip_id and asset_id:
                    referenced_clips[clip_id] = asset_id

        if referenced_assets:
            owned_assets = set((await db.execute(
                select(Asset.id).where(
                    Asset.project_id == script.project_id,
                    Asset.id.in_(referenced_assets),
                )
            )).scalars().all())
            if owned_assets != referenced_assets:
                raise HTTPException(status_code=400, detail="统一素材库包含不属于当前项目的素材")

        if referenced_clips:
            clip_rows = (await db.execute(
                select(Clip.id, Clip.asset_id).where(Clip.id.in_(set(referenced_clips)))
            )).all()
            actual_pairs = {clip_id: asset_id for clip_id, asset_id in clip_rows}
            if actual_pairs != referenced_clips:
                raise HTTPException(status_code=400, detail="统一素材引用的片段与素材关系无效")

        # 合并 visual_config
        current = script.visual_config or {}
        current.update(update_data["visual_config"])
        script.visual_config = current
        update_data.pop("visual_config")

    for key, value in update_data.items():
        setattr(script, key, value)

    script.needs_initial_save = False
    await db.flush()

    # 重新加载脚本（含 units 和 shots），避免异步序列化错误
    fresh_result = await db.execute(
        select(Script)
        .where(Script.id == script_id)
        .options(selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations))
    )
    script = fresh_result.scalar_one()
    return ScriptOut.model_validate(script)


@router.patch("/{script_id}/units")
async def batch_update_units(
    script_id: str,
    req: ScriptUnitBatchUpdate,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """批量更新分镜单元（保存脚本）。"""
    result = await db.execute(
        select(Script).join(Project, Script.project_id == Project.id).where(
            Script.id == script_id, accessible_project_condition(user.id)
        )
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    await require_project_access(script.project_id, user, db, "editor")

    selected_clip_ids = {
        clip_id
        for unit_data in req.units
        for clip_id in (
            [unit_data.get("selected_clip_id")]
            + [shot.get("selected_clip_id") for shot in unit_data.get("shots", [])]
        )
        if clip_id
    }
    if selected_clip_ids:
        allowed_clip_ids = set(
            (
                await db.execute(
                    select(Clip.id)
                    .join(Asset, Clip.asset_id == Asset.id)
                    .where(
                        Asset.project_id == script.project_id,
                        Clip.id.in_(selected_clip_ids),
                    )
                )
            ).scalars().all()
        )
        if allowed_clip_ids != selected_clip_ids:
            raise HTTPException(status_code=400, detail="脚本包含不属于当前项目的素材片段")

    # 获取现有 units
    units_result = await db.execute(
        select(ScriptUnit)
        .where(ScriptUnit.script_id == script_id)
        .options(selectinload(ScriptUnit.shots).selectinload(Shot.generations))
        .order_by(ScriptUnit.index)
    )
    existing_units = {u.id: u for u in units_result.scalars().all()}
    incoming_unit_ids = {
        str(unit_data.get("id"))
        for unit_data in req.units
        if unit_data.get("id") in existing_units
    }

    # 前端删除的脚本单元必须同步删除，不能在保存后重新出现。
    for removed_unit_id in set(existing_units) - incoming_unit_ids:
        removed_unit = existing_units[removed_unit_id]
        removed_shot_ids = [shot.id for shot in removed_unit.shots]
        if removed_shot_ids:
            await db.execute(delete(Candidate).where(Candidate.shot_id.in_(removed_shot_ids)))
        await db.execute(delete(Candidate).where(Candidate.unit_id == removed_unit_id))
        await db.delete(removed_unit)

    for idx, unit_data in enumerate(req.units):
        unit_id = unit_data.get("id")
        if unit_id and unit_id in existing_units:
            unit = existing_units[unit_id]
            unit_semantic_changed = any(
                unit_data.get(field, getattr(unit, field)) != getattr(unit, field)
                for field in ("outline", "text", "intent")
            )
            unit.index = idx
            unit.outline = unit_data.get("outline", unit.outline)
            unit.text = unit_data.get("text", unit.text)
            unit.intent = unit_data.get("intent", unit.intent)
            unit.status = unit_data.get("status", unit.status)
            unit.selected_clip_id = unit_data.get("selected_clip_id", unit.selected_clip_id)
            unit.duration = unit_data.get("duration", unit.duration)
            unit.notes = unit_data.get("notes", unit.notes)
            if unit_semantic_changed:
                retained_shot_ids = [shot.id for shot in unit.shots]
                if retained_shot_ids:
                    await db.execute(delete(Candidate).where(Candidate.shot_id.in_(retained_shot_ids)))
                await db.execute(delete(Candidate).where(Candidate.unit_id == unit.id))
                unit.candidates = []
                unit.status = "待编辑"

            # 更新 shots
            shots_data = unit_data.get("shots", [])
            existing_shots = {shot.id: shot for shot in unit.shots}
            incoming_shot_ids = {
                str(shot_data.get("id"))
                for shot_data in shots_data
                if shot_data.get("id") in existing_shots
            }
            for removed_shot_id in set(existing_shots) - incoming_shot_ids:
                await db.execute(delete(Candidate).where(Candidate.shot_id == removed_shot_id))
                await db.delete(existing_shots[removed_shot_id])
            for s_idx, shot_data in enumerate(shots_data):
                shot_id = shot_data.get("id")
                if shot_id and shot_id in existing_shots:
                    shot = existing_shots[shot_id]
                    if shot:
                        description_changed = (
                            shot_data.get("description", shot.description) != shot.description
                        )
                        shot.index = s_idx
                        shot.description = shot_data.get("description", shot.description)
                        shot.duration = shot_data.get("duration", shot.duration)
                        shot.notes = shot_data.get("notes", shot.notes)
                        shot.visual_mode = shot_data.get("visual_mode", shot.visual_mode)
                        shot.status = shot_data.get("status", shot.status)
                        incoming_clip_id = shot_data.get("selected_clip_id", shot.selected_clip_id)
                        shot.selected_clip_id = incoming_clip_id
                        if incoming_clip_id:
                            await db.execute(
                                update(ShotGeneration)
                                .where(ShotGeneration.shot_id == shot.id)
                                .values(is_selected=False)
                            )
                            shot.generated_image_path = None
                            shot.generated_video_path = None
                            shot.generated_effect_video_path = None
                            shot.visual_mode = "recommend"
                        if unit_semantic_changed or description_changed:
                            await db.execute(delete(Candidate).where(Candidate.shot_id == shot.id))
                            shot.candidates = []
                            shot.status = "待编辑"
                            shot.selected_clip_id = None
                            shot.generated_image_path = None
                            shot.generated_video_path = None
                            shot.generated_effect_video_path = None
                            await db.execute(
                                update(ShotGeneration)
                                .where(ShotGeneration.shot_id == shot.id)
                                .values(is_selected=False)
                            )
                else:
                    shot = Shot(
                        unit_id=unit.id,
                        index=s_idx,
                        description=shot_data.get("description", ""),
                        duration=shot_data.get("duration", 8),
                        notes=shot_data.get("notes", ""),
                        visual_mode=shot_data.get("visual_mode"),
                        status=shot_data.get("status", "待编辑"),
                        selected_clip_id=shot_data.get("selected_clip_id"),
                    )
                    db.add(shot)
        else:
            # 新建 unit
            unit = ScriptUnit(
                script_id=script_id,
                index=idx,
                outline=unit_data.get("outline", ""),
                text=unit_data.get("text", ""),
                intent=unit_data.get("intent", ""),
                duration=unit_data.get("duration", 8),
                notes=unit_data.get("notes", ""),
                status="待编辑",
            )
            db.add(unit)
            await db.flush()
            for s_idx, shot_data in enumerate(unit_data.get("shots", [])):
                shot = Shot(
                    unit_id=unit.id,
                    index=s_idx,
                    description=shot_data.get("description", ""),
                    duration=shot_data.get("duration", 8),
                    notes=shot_data.get("notes", ""),
                    visual_mode=shot_data.get("visual_mode"),
                    status="待编辑",
                )
                db.add(shot)

    await db.flush()

    # 重算指标
    units_result = await db.execute(
        select(ScriptUnit).where(ScriptUnit.script_id == script_id)
        .options(selectinload(ScriptUnit.shots).selectinload(Shot.generations))
        .order_by(ScriptUnit.index)
    )
    units = list(units_result.scalars().all())
    metrics = _compute_metrics(units)
    script.duration = metrics["duration"]
    script.unit_count = metrics["unit_count"]
    script.matched_count = metrics["matched_count"]
    script.progress = metrics["progress"]
    script.needs_initial_save = False

    await db.flush()
    return {"status": "ok", "progress": script.progress, "matched_count": script.matched_count}


@router.patch("/{script_id}/editor-state", response_model=ScriptOut)
async def save_script_editor_state(
    script_id: str,
    req: ScriptEditorSave,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """在同一数据库事务中保存标题、视觉配置、脚本句和全部分镜。"""
    await update_script(
        script_id,
        ScriptUpdate(title=req.title, visual_config=req.visual_config),
        user,
        db,
    )
    await batch_update_units(
        script_id,
        ScriptUnitBatchUpdate(units=req.units),
        user,
        db,
    )
    fresh = (
        await db.execute(
            select(Script)
            .where(Script.id == script_id)
            .options(
                selectinload(Script.units)
                .selectinload(ScriptUnit.shots)
                .selectinload(Shot.generations)
            )
        )
    ).scalar_one()
    return ScriptOut.model_validate(fresh)


@router.post(
    "/{script_id}/shots/{shot_id}/generations/{generation_id}/select",
    response_model=ShotOut,
)
async def select_shot_generation(
    script_id: str,
    shot_id: str,
    generation_id: str,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """显式采用一条历史生成结果；其他结果继续保留为可选历史。"""
    row = (
        await db.execute(
            select(ShotGeneration, Shot, Script)
            .join(Shot, Shot.id == ShotGeneration.shot_id)
            .join(ScriptUnit, ScriptUnit.id == Shot.unit_id)
            .join(Script, Script.id == ScriptUnit.script_id)
            .join(Project, Project.id == Script.project_id)
            .where(
                ShotGeneration.id == generation_id,
                Shot.id == shot_id,
                Script.id == script_id,
                accessible_project_condition(user.id),
            )
        )
    ).first()
    if not row:
        raise HTTPException(status_code=404, detail="生成结果不存在或不属于当前分镜")
    generation, shot, script = row
    await require_project_access(script.project_id, user, db, "editor")

    await db.execute(
        update(ShotGeneration)
        .where(ShotGeneration.shot_id == shot.id)
        .values(is_selected=False)
    )
    generation.is_selected = True
    shot.generated_image_path = generation.storage_path if generation.kind == "image" else None
    shot.generated_video_path = generation.storage_path if generation.kind == "video" else None
    shot.generated_effect_video_path = generation.storage_path if generation.kind == "effect" else None
    shot.visual_mode = {"image": "image", "video": "video", "effect": "sketch"}[generation.kind]
    shot.selected_clip_id = None
    shot.status = "已选择"
    await db.flush()

    fresh_shot = (
        await db.execute(
            select(Shot)
            .where(Shot.id == shot.id)
            .options(selectinload(Shot.generations))
        )
    ).scalar_one()
    return ShotOut.model_validate(fresh_shot)


@router.post("/{script_id}/match")
async def match_script_clips(
    script_id: str,
    req: MatchRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """为指定分镜、指定单元或全部分镜创建真实镜头匹配任务。"""
    if not settings.has_embedding:
        raise HTTPException(status_code=503, detail="语义匹配服务未配置或不可用")

    result = await db.execute(
        select(Script)
        .join(Project, Script.project_id == Project.id)
        .where(Script.id == script_id, accessible_project_condition(user.id))
        .options(selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations))
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    await require_project_access(script.project_id, user, db, "editor")

    units = [unit for unit in script.units if not req.unit_id or unit.id == req.unit_id]
    if req.unit_id and not units:
        raise HTTPException(status_code=404, detail="脚本单元不存在")
    available_shot_ids = [shot.id for unit in units for shot in (unit.shots or [])]
    requested_shot_ids = list(dict.fromkeys(req.shot_ids))
    if requested_shot_ids:
        invalid_ids = set(requested_shot_ids) - set(available_shot_ids)
        if invalid_ids:
            raise HTTPException(status_code=404, detail="部分分镜不存在或不属于当前脚本")
        shot_ids = requested_shot_ids
    else:
        shot_ids = available_shot_ids
    if not shot_ids:
        raise HTTPException(status_code=400, detail="脚本中没有可匹配的分镜")

    job = Job(
        project_id=script.project_id,
        name=f"镜头匹配：{script.title}",
        type="镜头匹配",
        status="QUEUED",
        stage="等待匹配",
        stages=["生成语义", "检索素材", "写入候选"],
        stage_index=0,
        progress=0,
        eta="预计 1-2 分钟",
        detail=f"准备匹配 {len(shot_ids)} 个分镜",
        result=None,
        result_detail="",
        payload={"script_id": script.id, "shot_ids": shot_ids},
        is_simulated=False,
    )
    db.add(job)
    await db.commit()

    from app.tasks.match_tasks import match_script_shots
    from app.services.task_dispatch import dispatch_project_task
    try:
        task = await dispatch_project_task(
            db, script.project_id, match_script_shots, script.id, shot_ids, job.id
        )
        job.celery_task_id = task.id
        await db.commit()
    except Exception as exc:
        job.status = "FAILED"
        job.stage = "任务派发失败"
        job.detail = f"镜头匹配任务派发失败：{str(exc)[:200]}"
        job.result = None
        job.result_detail = "请确认 Redis 和 Worker 正常运行后重试"
        await db.commit()
    return {"job_id": job.id, "message": "镜头匹配任务已创建"}


@router.post("/{script_id}/complete-shots")
async def complete_script_shots(
    script_id: str,
    req: CompleteShotsRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """使用一次显式线上文本调用完善用户选中的分镜描述。"""
    if not settings.has_text_generation:
        raise HTTPException(status_code=503, detail="AI 文本生成服务未配置或不可用")
    script = (
        await db.execute(
            select(Script)
            .join(Project, Script.project_id == Project.id)
            .where(Script.id == script_id, accessible_project_condition(user.id))
        )
    ).scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    await require_project_access(script.project_id, user, db, "editor")

    shot_ids = list(dict.fromkeys(req.shot_ids))
    owned_shot_ids = set(
        (
            await db.execute(
                select(Shot.id)
                .join(ScriptUnit, Shot.unit_id == ScriptUnit.id)
                .where(ScriptUnit.script_id == script_id, Shot.id.in_(shot_ids))
            )
        ).scalars().all()
    )
    if owned_shot_ids != set(shot_ids):
        raise HTTPException(status_code=404, detail="部分分镜不存在或不属于当前脚本")

    job = Job(
        project_id=script.project_id,
        name=f"AI 完善分镜：{script.title}",
        type="AI 分镜完善",
        status="QUEUED",
        stage="等待完善",
        stages=["读取脚本", "完善描述", "写入分镜"],
        stage_index=0,
        progress=0,
        eta="预计 30-60 秒",
        detail=f"准备完善 {len(shot_ids)} 个分镜",
        result=None,
        result_detail="",
        payload={"script_id": script.id, "shot_ids": shot_ids},
        is_simulated=False,
    )
    db.add(job)
    await db.commit()

    from app.tasks.script_tasks import complete_shot_descriptions
    from app.services.task_dispatch import dispatch_project_task
    try:
        task = await dispatch_project_task(
            db, script.project_id, complete_shot_descriptions, script.id, shot_ids, job.id
        )
        job.celery_task_id = task.id
        await db.commit()
    except Exception as exc:
        job.status = "FAILED"
        job.stage = "任务派发失败"
        job.detail = f"AI 分镜完善任务派发失败：{str(exc)[:200]}"
        job.result_detail = "请确认 Redis 和 Worker 正常运行后重试"
        await db.commit()
    return {"job_id": job.id, "message": "AI 分镜完善任务已创建"}


@router.post("/{script_id}/rough-cut", response_model=dict)
async def create_rough_cut(
    script_id: str,
    req: RoughCutRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """从脚本生成 AI 初剪项目。

    对应详设 3.3 createProjectFromScript：
    复制脚本 → 创建项目 → 派发初剪任务
    """
    if not settings.has_embedding:
        raise HTTPException(status_code=503, detail="AI 初剪所需的语义向量服务未配置或不可用")
    result = await db.execute(
        select(Script).join(Project, Script.project_id == Project.id).where(
            Script.id == script_id, accessible_project_condition(user.id)
        )
        .options(selectinload(Script.units).selectinload(ScriptUnit.shots).selectinload(Shot.generations))
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    await require_project_access(script.project_id, user, db, "editor")

    if req.visual_config:
        script.visual_config = req.visual_config
    visual_config = script.visual_config or {}

    # 同一来源脚本只维护一个剪辑项目；进行中的版本直接返回，防止双击重复创建。
    existing_project = (
        await db.execute(
            select(Project)
            .where(
                Project.owner_id == user.id,
                Project.source_script_id == script.id,
            )
            .order_by(Project.created_at)
        )
    ).scalars().first()
    if existing_project:
        running_version = (
            await db.execute(
                select(CutVersion)
                .where(
                    CutVersion.project_id == existing_project.id,
                    CutVersion.status == "AI 初剪生成中",
                )
                .order_by(desc(CutVersion.created_at))
            )
        ).scalars().first()
        existing_job = (
            await db.execute(
                select(Job)
                .where(
                    Job.project_id == existing_project.id,
                    Job.type == "AI 初剪生成",
                    Job.status.in_(["QUEUED", "RUNNING"]),
                )
                .order_by(desc(Job.created_at))
            )
        ).scalars().first()
        if running_version and running_version.snapshot_script_id and existing_job:
            return {
                "project_id": existing_project.id,
                "script_id": running_version.snapshot_script_id,
                "version_id": running_version.id,
                "job_id": existing_job.id,
                "message": "该脚本已有正在生成的剪辑版本",
            }

    if existing_project:
        ai_project = existing_project
        ai_project.status = "AI 初剪生成中"
        ai_project.visual_style = visual_config.get("stylePreset", "brand-tech")
        ai_project.aspect_ratio = visual_config.get("aspectRatio", "16:9")
    else:
        ai_project = Project(
            owner_id=user.id,
            name=f"{script.title} · 剪辑项目",
            initial=script.title[0].upper() if script.title else "A",
            type="AI 剪辑项目",
            description=f"基于脚本「{script.title}」创建的多版本剪辑项目",
            status="AI 初剪生成中",
            variant="robot",
            source_script_id=script.id,
            visual_style=visual_config.get("stylePreset", "brand-tech"),
            aspect_ratio=visual_config.get("aspectRatio", "16:9"),
            ai_shot_count=0,
            is_new=True,
        )
        db.add(ai_project)
        await db.flush()

    version_count = (
        await db.execute(
            select(func.count()).select_from(CutVersion).where(CutVersion.project_id == ai_project.id)
        )
    ).scalar() or 0

    # 复制脚本
    import copy
    copied_script = Script(
        project_id=ai_project.id,
        title=f"{script.title}（初剪副本）",
        summary=f"正在为 {script.unit_count} 个脚本单元装配真实素材初剪。",
        status="初剪生成中",
        duration=script.duration,
        unit_count=script.unit_count,
        matched_count=script.matched_count,
        progress=script.progress,
        variant=script.variant,
        format=script.format,
        visual_config=copy.deepcopy(script.visual_config),
        source_type="ai",
        is_snapshot=True,
        parent_script_id=script.id,
        needs_initial_save=False,
    )
    db.add(copied_script)
    await db.flush()

    cut_version = CutVersion(
        project_id=ai_project.id,
        snapshot_script_id=copied_script.id,
        name=f"V{version_count + 1}",
        status="AI 初剪生成中",
        aspect_ratio=visual_config.get("aspectRatio", "16:9"),
        timeline=[],
        timeline_version=0,
    )
    db.add(cut_version)
    await db.flush()

    # 复制 units 和 shots
    for unit in script.units:
        copied_unit = ScriptUnit(
            script_id=copied_script.id,
            source_unit_id=unit.id,
            index=unit.index,
            outline=unit.outline,
            text=unit.text,
            intent=unit.intent,
            status=unit.status,
            selected_clip_id=unit.selected_clip_id,
            duration=unit.duration,
            notes=unit.notes,
        )
        db.add(copied_unit)
        await db.flush()
        for shot in (unit.shots or []):
            copied_shot = Shot(
                unit_id=copied_unit.id,
                source_shot_id=shot.id,
                index=shot.index,
                description=shot.description,
                status="已选择" if shot.selected_clip_id else "待编辑",
                selected_clip_id=shot.selected_clip_id,
                visual_mode=shot.visual_mode,
                duration=shot.duration,
                notes=shot.notes,
                generated_image_path=shot.generated_image_path,
                generated_video_path=shot.generated_video_path,
                generated_effect_video_path=shot.generated_effect_video_path,
            )
            db.add(copied_shot)

    # 创建初剪任务
    job = Job(
        project_id=ai_project.id,
        name=f"AI 初剪：{script.title}",
        type="AI 初剪生成",
        status="QUEUED",
        stage="创建项目",
        stages=["创建项目", "锁定统一视觉", "匹配真实素材与识别缺失镜头", "装配多轨初剪"],
        stage_index=0,
        progress=0,
        eta="预计 1-2 分钟",
        detail="AI 初剪任务已进入队列",
        result=None,
        result_detail="",
        payload={
            "project_id": ai_project.id,
            "script_id": copied_script.id,
            "source_project_id": script.project_id,
            "cut_version_id": cut_version.id,
        },
        is_simulated=False,
    )
    db.add(job)
    await db.commit()

    # 派发 Celery 任务
    from app.tasks.roughcut_tasks import run_rough_cut
    from app.services.task_dispatch import dispatch_project_task
    try:
        task = await dispatch_project_task(
            db,
            ai_project.id,
            run_rough_cut,
            job.id,
            ai_project.id,
            copied_script.id,
            script.project_id,
            cut_version.id,
        )
        job.celery_task_id = task.id
        await db.commit()
    except Exception as exc:
        job.status = "FAILED"
        job.stage = "任务派发失败"
        job.detail = f"AI 初剪任务派发失败：{str(exc)[:200]}"
        job.result = None
        job.result_detail = "请确认 Redis 和 Worker 正常运行后重试"
        await db.commit()

    return {
        "project_id": ai_project.id,
        "script_id": copied_script.id,
        "version_id": cut_version.id,
        "job_id": job.id,
        "message": "AI 初剪任务已创建",
    }


@router.delete("/{script_id}", status_code=204)
async def delete_script(
    script_id: str,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """仅在脚本没有剪辑项目、快照、媒体索引或运行任务依赖时删除。"""
    result = await db.execute(
        select(Script).join(Project, Script.project_id == Project.id).where(
            Script.id == script_id, accessible_project_condition(user.id)
        )
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    await require_project_access(script.project_id, user, db, "editor")
    if script.is_snapshot:
        raise HTTPException(status_code=400, detail="初剪快照由剪辑项目管理，不能单独删除")

    from app.services.deletion import script_delete_impact

    impact = await script_delete_impact(db, script)
    if not impact["can_delete"]:
        raise HTTPException(status_code=409, detail="；".join(impact["blockers"]))

    project = await db.get(Project, script.project_id)
    unit_ids = list(
        (await db.execute(select(ScriptUnit.id).where(ScriptUnit.script_id == script.id))).scalars().all()
    )
    shot_ids: list[str] = []
    generated_paths: set[str] = set()
    if unit_ids:
        shot_rows = (
            await db.execute(
                select(
                    Shot.id,
                    Shot.generated_image_path,
                    Shot.generated_video_path,
                    Shot.generated_effect_video_path,
                ).where(Shot.unit_id.in_(unit_ids))
            )
        ).all()
        shot_ids = [row[0] for row in shot_rows]
        for row in shot_rows:
            generated_paths.update(path for path in row[1:] if path)
        await db.execute(delete(Candidate).where(Candidate.unit_id.in_(unit_ids)))
    if shot_ids:
        await db.execute(delete(Candidate).where(Candidate.shot_id.in_(shot_ids)))
    await db.delete(script)
    if project:
        remaining_script_count = (
            await db.execute(
                select(func.count())
                .select_from(Script)
                .where(Script.project_id == project.id, Script.id != script.id)
            )
        ).scalar() or 0
        project.script_count = remaining_script_count
    await db.commit()

    from app.utils.storage import delete_file

    for path in generated_paths:
        delete_file(path)


@router.get("/{script_id}/delete-impact")
async def get_script_delete_impact(
    script_id: str,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Script)
        .join(Project, Script.project_id == Project.id)
        .where(
            Script.id == script_id,
            Script.is_snapshot.is_(False),
            accessible_project_condition(user.id),
        )
        .options(selectinload(Script.units))
    )
    script = result.scalar_one_or_none()
    if not script:
        raise HTTPException(status_code=404, detail="脚本不存在")
    await require_project_access(script.project_id, user, db, "editor")
    from app.services.deletion import script_delete_impact

    return await script_delete_impact(db, script)
